#!/usr/bin/env python3
"""Generate ISP Project Management raw data for Power BI + web dashboard."""

import json
import random
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

from faker import Faker
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

random.seed(42)
fake = Faker("en_IN")
Faker.seed(42)

OUT_DIR = Path(__file__).parent
XLSX_PATH = OUT_DIR / "Isp_Pm_RawData_PowerBI.xlsx"
JSON_PATH = OUT_DIR / "dashboard-data.json"

CLIENTS = [
    ("CL001", "Hitachi Payment Services", "BFSI", "Enterprise", "Mumbai"),
    ("CL002", "Axis Bank", "BFSI", "Enterprise", "Mumbai"),
    ("CL003", "State Bank of India", "BFSI", "Enterprise", "Mumbai"),
    ("CL004", "Bajaj Finance", "BFSI", "Enterprise", "Pune"),
    ("CL005", "HDFC Bank", "BFSI", "Enterprise", "Mumbai"),
    ("CL006", "ICICI Bank", "BFSI", "Enterprise", "Mumbai"),
    ("CL007", "Tata Motors", "Manufacturing", "Large", "Pune"),
    ("CL008", "Reliance Retail", "Retail", "Enterprise", "Mumbai"),
]

PROJECT_TYPES = ["SD-WAN Rollout", "MPLS Backbone", "VSAT Backup", "ATM Connectivity", "Branch Link Upgrade"]
REGIONS = ["West", "North", "South", "East", "Central"]
CITIES = {
    "West": ["Mumbai", "Pune", "Ahmedabad", "Surat"],
    "North": ["Delhi", "Noida", "Gurgaon", "Jaipur"],
    "South": ["Bengaluru", "Chennai", "Hyderabad", "Kochi"],
    "East": ["Kolkata", "Bhubaneswar", "Patna"],
    "Central": ["Bhopal", "Indore", "Nagpur"],
}
SITE_STATUSES = ["Survey Pending", "ROW In Progress", "Installation Scheduled", "Installation Done", "Testing", "Live", "Pending Acceptance"]
TASK_PHASES = ["Initiation", "Survey", "ROW/Permits", "Procurement", "Installation", "Testing", "Go-Live", "Hypercare"]
TASK_STATUSES = ["Not Started", "In Progress", "Blocked", "Completed", "Delayed"]
RAG = ["Green", "Amber", "Red"]
RAID_TYPES = ["Risk", "Issue", "Assumption", "Dependency"]
SEVERITY = ["Low", "Medium", "High", "Critical"]
SLA_METRICS = ["Site Survey TAT", "ROW Clearance TAT", "Installation TAT", "Go-Live TAT", "Fault Resolution", "Uptime SLA"]
COST_CATEGORIES = ["CapEx Hardware", "Last Mile", "Labor", "ROW Charges", "Logistics", "PMO Overhead"]
RESOURCE_ROLES = ["Project Manager", "Site Engineer", "Network Engineer", "ROW Coordinator", "QA Lead", "Delivery Manager"]
CERTS = ["PMP", "PRINCE2", "CCNA", "CCNP", "CSM", "ITIL"]

TODAY = date(2026, 6, 29)


def fmt(d):
    return d.strftime("%Y-%m-%d") if d else ""


def rag_from_pct(pct, delay=0):
    if pct >= 85 and delay <= 3:
        return "Green"
    if pct >= 60 or delay <= 10:
        return "Amber"
    return "Red"


def generate():
    clients = []
    for cid, name, sector, tier, hq in CLIENTS:
        clients.append({
            "ClientID": cid, "ClientName": name, "Sector": sector, "Tier": tier,
            "HQCity": hq, "AccountOwner": fake.name(), "ContractValue_Cr": round(random.uniform(2.5, 45), 2),
            "ActiveProjects": 0,
        })

    projects = []
    sites = []
    tasks = []
    raid = []
    resources = []
    budget = []
    sla = []
    weekly = []

    pm_pool = [fake.name() for _ in range(12)]
    res_id = 1
    for i in range(65):
        resources.append({
            "ResourceID": f"RES{res_id:03d}",
            "Name": fake.name(),
            "Role": random.choice(RESOURCE_ROLES),
            "Region": random.choice(REGIONS),
            "BillRate_INR": random.choice([2500, 3200, 4500, 5500, 6800]),
            "Utilization_Pct": random.randint(55, 98),
            "Certification": random.choice(CERTS),
            "Experience_Yrs": random.randint(3, 18),
            "Status": random.choice(["Allocated", "Bench", "On Leave"]),
        })
        res_id += 1

    site_seq = 1
    task_seq = 1
    raid_seq = 1
    budget_seq = 1
    sla_seq = 1
    weekly_seq = 1

    for pi, (cid, cname, *_ ) in enumerate(CLIENTS):
        n_projects = random.randint(4, 6)
        for pj in range(n_projects):
            pid = f"PRJ{len(projects)+1:03d}"
            ptype = random.choice(PROJECT_TYPES)
            region = random.choice(REGIONS)
            start = TODAY - timedelta(days=random.randint(120, 400))
            planned_end = start + timedelta(days=random.randint(180, 360))
            pct = random.randint(35, 95)
            delay = random.randint(0, 18)
            budget_cr = round(random.uniform(1.2, 12), 2)
            actual_cr = round(budget_cr * random.uniform(0.55, 1.08), 2)
            rag = rag_from_pct(pct, delay)
            pm = random.choice(pm_pool)
            n_sites = random.randint(60, 120)

            projects.append({
                "ProjectID": pid, "ClientID": cid, "ClientName": cname,
                "ProjectName": f"{cname.split()[0]} {ptype} Phase {pj+1}",
                "ProjectType": ptype, "Region": region, "PMOwner": pm,
                "StartDate": fmt(start), "PlannedEndDate": fmt(planned_end),
                "Budget_Cr": budget_cr, "ActualSpend_Cr": actual_cr,
                "PctComplete": pct, "DelayDays": delay, "RAGStatus": rag,
                "TotalSites": n_sites, "SitesLive": int(n_sites * pct / 100),
                "OpenIssues": random.randint(0, 8), "OpenRisks": random.randint(0, 5),
            })

            for si in range(n_sites):
                sid = f"ST{site_seq:05d}"
                site_seq += 1
                city = random.choice(CITIES[region])
                status = random.choices(
                    SITE_STATUSES,
                    weights=[8, 12, 15, 18, 10, 28, 9],
                )[0]
                survey = start + timedelta(days=random.randint(5, 30))
                install_planned = survey + timedelta(days=random.randint(20, 60))
                install_actual = install_planned + timedelta(days=random.randint(-5, 10)) if status not in ["Survey Pending", "ROW In Progress"] else None
                golive_planned = install_planned + timedelta(days=random.randint(10, 30))
                golive_actual = golive_planned + timedelta(days=random.randint(-3, 12)) if status in ["Live", "Pending Acceptance"] else None

                if status in ["Live", "Pending Acceptance"] and install_actual:
                    delay_days = max(0, (install_actual - install_planned).days)
                elif status in ["Installation Done", "Testing"]:
                    delay_days = max(0, min(12, (TODAY - install_planned).days)) if TODAY > install_planned else 0
                else:
                    delay_days = random.randint(0, 8)

                sla_status = "Met" if delay_days <= 7 else ("Marginal" if delay_days <= 14 else "Breached")

                sites.append({
                    "SiteID": sid, "ProjectID": pid, "ClientID": cid, "ClientName": cname,
                    "SiteName": f"{city} Site-{si+1:03d}", "City": city, "Region": region,
                    "Bandwidth_Mbps": random.choice([10, 20, 50, 100, 200]),
                    "SiteStatus": status, "SurveyDate": fmt(survey),
                    "InstallPlanned": fmt(install_planned),
                    "InstallActual": fmt(install_actual),
                    "GoLivePlanned": fmt(golive_planned),
                    "GoLiveActual": fmt(golive_actual),
                    "DelayDays": delay_days, "SLAStatus": sla_status,
                    "Engineer": random.choice(resources)["Name"],
                    "SiteCost_INR": random.randint(45000, 280000),
                })

            for phase in TASK_PHASES:
                t_pct = min(100, max(0, pct + random.randint(-15, 15)))
                t_status = "Completed" if t_pct == 100 else random.choice(TASK_STATUSES)
                tasks.append({
                    "TaskID": f"TSK{task_seq:04d}", "ProjectID": pid, "ClientID": cid,
                    "WBSPhase": phase, "TaskName": f"{phase} - {ptype}",
                    "Owner": pm if phase in ["Initiation", "Go-Live"] else random.choice(resources)["Name"],
                    "PlannedStart": fmt(start + timedelta(days=TASK_PHASES.index(phase) * 25)),
                    "PlannedEnd": fmt(start + timedelta(days=(TASK_PHASES.index(phase) + 1) * 25)),
                    "PctComplete": t_pct, "Status": t_status,
                    "DelayDays": random.randint(0, 12) if t_status == "Delayed" else 0,
                })
                task_seq += 1

            for _ in range(random.randint(3, 6)):
                raid.append({
                    "RAIDID": f"RD{raid_seq:03d}", "ProjectID": pid, "ClientID": cid,
                    "Type": random.choice(RAID_TYPES), "Title": fake.sentence(nb_words=6)[:-1],
                    "Severity": random.choice(SEVERITY), "Owner": random.choice(pm_pool),
                    "Status": random.choice(["Open", "Mitigated", "Closed", "Escalated"]),
                    "Impact": random.choice(["Schedule", "Cost", "Quality", "Scope"]),
                    "DueDate": fmt(TODAY + timedelta(days=random.randint(-10, 45))),
                    "DaysOpen": random.randint(1, 60),
                })
                raid_seq += 1

            for month_offset in range(6):
                m_start = (TODAY.replace(day=1) - timedelta(days=30 * month_offset)).replace(day=1)
                for cat in random.sample(COST_CATEGORIES, k=random.randint(3, 5)):
                    planned = random.randint(800000, 4500000)
                    actual = int(planned * random.uniform(0.82, 1.12))
                    budget.append({
                        "BudgetID": f"BG{budget_seq:04d}", "ProjectID": pid, "ClientID": cid,
                        "Month": m_start.strftime("%Y-%m"), "Category": cat,
                        "Planned_INR": planned, "Actual_INR": actual,
                        "Variance_Pct": round((actual - planned) / planned * 100, 1),
                    })
                    budget_seq += 1

            for metric in random.sample(SLA_METRICS, k=random.randint(3, 5)):
                target = random.choice([95, 96, 97, 98, 99])
                achieved = round(random.uniform(target - 4, target + 1.5), 1)
                status = "Met" if achieved >= target else ("Marginal" if achieved >= target - 2 else "Breached")
                sla.append({
                    "SLAID": f"SL{sla_seq:03d}", "ProjectID": pid, "ClientID": cid,
                    "Metric": metric, "Target_Pct": target, "Achieved_Pct": achieved,
                    "Status": status, "Penalty_INR": 0 if status == "Met" else random.randint(50000, 500000),
                    "ReviewPeriod": "Q2 FY26",
                })
                sla_seq += 1

            for w in range(12):
                wk = TODAY - timedelta(weeks=w)
                weekly.append({
                    "LogID": f"WL{weekly_seq:04d}", "ProjectID": pid, "ClientID": cid,
                    "WeekEnding": fmt(wk), "RAGStatus": random.choice(RAG),
                    "PctComplete": max(10, pct - w * random.randint(1, 4)),
                    "Narrative": fake.paragraph(nb_sentences=2),
                    "OpenIssues": random.randint(0, 6),
                    "SitesLive_Week": random.randint(0, 8),
                })
                weekly_seq += 1

    for c in clients:
        c["ActiveProjects"] = sum(1 for p in projects if p["ClientID"] == c["ClientID"])

    dates = []
    d = date(2024, 1, 1)
    end = date(2026, 12, 31)
    while d <= end:
        fy = d.year if d.month >= 4 else d.year - 1
        q = (d.month - 1) // 3 + 1
        dates.append({
            "Date": fmt(d), "Year": d.year, "Month": d.month, "MonthName": d.strftime("%B"),
            "Quarter": f"Q{q}", "FY": f"FY{fy}-{str(fy+1)[-2:]}",
            "WeekNum": d.isocalendar()[1], "IsWeekend": d.weekday() >= 5,
        })
        d += timedelta(days=1)

    return {
        "Client_Master": clients,
        "Projects": projects,
        "Sites_Milestones": sites,
        "Tasks_WBS": tasks,
        "RAID_Log": raid,
        "Resources": resources,
        "Budget_Costs": budget,
        "SLA_Governance": sla,
        "Weekly_Status_Log": weekly,
        "Date_Dimension": dates,
    }


def write_xlsx(data):
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1B3A5C")

    for sheet_name, rows in data.items():
        ws = wb.create_sheet(sheet_name)
        if not rows:
            continue
        headers = list(rows[0].keys())
        for ci, h in enumerate(headers, 1):
            cell = ws.cell(1, ci, h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for ri, row in enumerate(rows, 2):
            for ci, h in enumerate(headers, 1):
                ws.cell(ri, ci, row[h])
        for ci, h in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(ci)].width = min(28, max(12, len(h) + 4))
        ws.freeze_panes = "A2"

    wb.save(XLSX_PATH)


def aggregate_for_dashboard(data):
    projects = data["Projects"]
    sites = data["Sites_Milestones"]
    tasks = data["Tasks_WBS"]
    raid = data["RAID_Log"]
    resources = data["Resources"]
    budget = data["Budget_Costs"]
    sla = data["SLA_Governance"]
    clients = data["Client_Master"]

    def count_by(items, key):
        c = defaultdict(int)
        for x in items:
            c[x[key]] += 1
        return dict(c)

    def sum_by(items, group, val):
        c = defaultdict(float)
        for x in items:
            c[x[group]] += float(x[val])
        return {k: round(v, 2) for k, v in c.items()}

    total_budget = sum(p["Budget_Cr"] for p in projects)
    total_actual = sum(p["ActualSpend_Cr"] for p in projects)
    live_sites = sum(1 for s in sites if s["SiteStatus"] == "Live")
    breached_sla = sum(1 for s in sites if s["SLAStatus"] == "Breached")

    return {
        "meta": {
            "generated": datetime.now().isoformat(),
            "totalProjects": len(projects),
            "totalSites": len(sites),
            "totalTasks": len(tasks),
            "asOfDate": fmt(TODAY),
        },
        "kpis": {
            "activeProjects": len(projects),
            "totalSites": len(sites),
            "sitesLive": live_sites,
            "avgCompletion": round(sum(p["PctComplete"] for p in projects) / len(projects), 1),
            "budgetCr": round(total_budget, 2),
            "actualSpendCr": round(total_actual, 2),
            "budgetVariancePct": round((total_actual - total_budget) / total_budget * 100, 1),
            "openIssues": sum(1 for r in raid if r["Type"] == "Issue" and r["Status"] in ["Open", "Escalated"]),
            "openRisks": sum(1 for r in raid if r["Type"] == "Risk" and r["Status"] in ["Open", "Escalated"]),
            "slaBreachSites": breached_sla,
            "slaBreachPct": round(breached_sla / len(sites) * 100, 1),
            "avgUtilization": round(sum(r["Utilization_Pct"] for r in resources) / len(resources), 1),
        },
        "projects": projects,
        "clients": clients,
        "sitesByStatus": count_by(sites, "SiteStatus"),
        "sitesByRegion": count_by(sites, "Region"),
        "sitesBySLA": count_by(sites, "SLAStatus"),
        "sitesByClient": count_by(sites, "ClientName"),
        "ragDistribution": count_by(projects, "RAGStatus"),
        "projectTypeMix": count_by(projects, "ProjectType"),
        "tasksByStatus": count_by(tasks, "Status"),
        "tasksByPhase": count_by(tasks, "WBSPhase"),
        "raidByType": count_by(raid, "Type"),
        "raidBySeverity": count_by(raid, "Severity"),
        "raidOpen": [r for r in raid if r["Status"] in ["Open", "Escalated"]][:40],
        "slaByStatus": count_by(sla, "Status"),
        "slaMetrics": sla,
        "budgetByCategory": sum_by(budget, "Category", "Actual_INR"),
        "budgetMonthly": defaultdict(lambda: {"planned": 0, "actual": 0}),
        "resourcesByRole": count_by(resources, "Role"),
        "resourceUtilization": [{"name": r["Name"], "role": r["Role"], "util": r["Utilization_Pct"], "region": r["Region"]} for r in sorted(resources, key=lambda x: -x["Utilization_Pct"])[:20]],
        "weeklyTrend": [],
        "siteSample": sites[:200],
    }


def main():
    data = generate()
    write_xlsx(data)
    agg = aggregate_for_dashboard(data)

    # monthly budget trend
    monthly = defaultdict(lambda: {"planned": 0, "actual": 0})
    for b in data["Budget_Costs"]:
        monthly[b["Month"]]["planned"] += b["Planned_INR"]
        monthly[b["Month"]]["actual"] += b["Actual_INR"]
    agg["budgetMonthly"] = [
        {"month": m, "planned": round(v["planned"] / 1e7, 2), "actual": round(v["actual"] / 1e7, 2)}
        for m, v in sorted(monthly.items())
    ][-12:]

    # weekly completion trend
    wk = defaultdict(list)
    for w in data["Weekly_Status_Log"]:
        wk[w["WeekEnding"]].append(w["PctComplete"])
    agg["weeklyTrend"] = [
        {"week": k, "avgPct": round(sum(v) / len(v), 1)}
        for k, v in sorted(wk.items())
    ][-16:]

    with open(JSON_PATH, "w") as f:
        json.dump(agg, f, indent=2)

    total_rows = sum(len(v) for v in data.values())
    print(f"Excel: {XLSX_PATH} ({XLSX_PATH.stat().st_size // 1024} KB)")
    print(f"JSON:  {JSON_PATH} ({JSON_PATH.stat().st_size // 1024} KB)")
    print(f"Tables: {len(data)}, Total rows: {total_rows}")


if __name__ == "__main__":
    main()
